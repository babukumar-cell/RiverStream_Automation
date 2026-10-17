import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

output_dir = os.path.join(os.path.dirname(__file__), 'Test Cases', 'Module_Wise_Test_Cases')
os.makedirs(output_dir, exist_ok=True)

# Remove old files if they exist
for old_file in [
    '08_Email_Flow_Riverstone_and_CorpDesign.xlsx',
    '09_Email_Flow_Hightower.xlsx'
]:
    p = os.path.join(output_dir, old_file)
    if os.path.exists(p):
        os.remove(p)

HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
TITLE_FONT = Font(name='Calibri', size=14, bold=True, color='1F4E79')
SUBTITLE_FONT = Font(name='Calibri', size=10, italic=True, color='595959')
DATA_FONT = Font(name='Calibri', size=10)
PASS_FILL = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
READY_FILL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

HEADERS = [
    'Test Case ID',
    'Module / Feature',
    'Test Case Description',
    'Pre-conditions',
    'Steps to Reproduce',
    'Test Data',
    'Expected Result',
    'Automation Status',
    'Execution Result'
]

def format_sheet(ws, title, cases):
    ws.merge_cells('A1:I1')
    ws['A1'] = title
    ws['A1'].font = TITLE_FONT
    ws['A1'].alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:I2')
    ws['A2'] = f"River Stream Test Automation Framework • Organization: Riverstone • Total Cases: {len(cases)}"
    ws['A2'].font = SUBTITLE_FONT
    ws['A2'].alignment = Alignment(vertical='center')
    ws.row_dimensions[2].height = 18

    ws.row_dimensions[3].height = 8

    header_row = 4
    ws.row_dimensions[header_row].height = 24
    for col_idx, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = THIN_BORDER

    current_row = 5
    for case in cases:
        ws.row_dimensions[current_row].height = 42
        row_values = [
            case.get('id', ''),
            case.get('module', ''),
            case.get('desc', ''),
            case.get('pre', 'User logged in as Riverstone Admin'),
            case.get('steps', ''),
            case.get('data', '-'),
            case.get('expected', ''),
            case.get('auto_status', 'Automated'),
            case.get('exec_result', 'READY')
        ]
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            
            if col_idx in [1, 8, 9]:
                cell.alignment = Alignment(horizontal='center', vertical='top')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                
            if col_idx == 9 and val == 'PASSED':
                cell.fill = PASS_FILL
            elif col_idx == 9 and val == 'READY':
                cell.fill = READY_FILL

        current_row += 1

    col_widths = {
        'A': 16, 'B': 22, 'C': 35, 'D': 25, 'E': 42, 'F': 20, 'G': 42, 'H': 18, 'I': 16
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = 'A5'

with open('test_cases_extracted.json', 'r', encoding='utf-8') as f:
    raw_extracted = json.load(f)

# Load Riverstone raw workbook
riverstone_wb_path = os.path.join('Test Cases', 'River Stream Test cases.xlsx')
raw_wb = openpyxl.load_workbook(riverstone_wb_path, data_only=True)

modules_data = {
    '01_Authentication': [],
    '02_Dashboard_and_Left_Navigation': [],
    '03_Upload_Document': [],
    '04_Documents_Library': [],
    '05_PO_ACK_Recon': [],
    '06_Document_Review': [],
    '07_Settings_and_Riverstone_Organization': [],
    '08_Email_Flow_Riverstone_Organization': []
}

# 1. Authentication
auth_raw = raw_extracted.get('Sign-In page', [])
for r in auth_raw:
    if r and len(r) > 1 and str(r[0]).startswith('TC'):
        modules_data['01_Authentication'].append({
            'id': f"TC_AUTH_{str(r[0]).replace('TC_', '').zfill(3)}",
            'module': 'Authentication',
            'desc': str(r[1]).strip() if r[1] else '',
            'pre': 'Open River Stream URL (https://riverstream-qa.azurewebsites.net/)',
            'steps': str(r[2]).strip() if len(r) > 2 and r[2] else '',
            'data': str(r[3]).strip() if len(r) > 3 and r[3] else '',
            'expected': str(r[4]).strip() if len(r) > 4 and r[4] else '',
            'auto_status': 'Automated',
            'exec_result': 'PASSED'
        })

# 2. Left side & Dashboard & Documents & Settings split
left_raw = raw_extracted.get('Left side fields', [])
current_section = 'Dashboard'
for r in left_raw:
    cleaned = [str(c).strip() for c in r if c and str(c).strip() and str(c) != 'None']
    if not cleaned:
        continue
    first_val = str(r[0]).strip() if r[0] else ''
    if not first_val.startswith('TC') and not first_val.isdigit():
        if 'dashboard' in cleaned[0].lower():
            current_section = 'Dashboard'
        elif 'document' in cleaned[0].lower():
            current_section = 'Documents Library'
        elif 'setting' in cleaned[0].lower() or 'organization' in cleaned[0].lower():
            current_section = 'Settings & Org'
        elif 'client' in cleaned[0].lower():
            current_section = 'Clients'
        elif 'analytics' in cleaned[0].lower():
            current_section = 'Analytics'
        continue

    item = {
        'id': f"TC_{current_section[:4].upper()}_{first_val.replace('TC_', '').zfill(3)}",
        'module': current_section,
        'desc': str(r[1]).strip() if len(r) > 1 and r[1] else '',
        'pre': 'Logged in to River Stream Dashboard (Riverstone Org)',
        'steps': str(r[2]).strip() if len(r) > 2 and r[2] else '',
        'data': str(r[3]).strip() if len(r) > 3 and r[3] else '',
        'expected': str(r[4]).strip() if len(r) > 4 and r[4] else '',
        'auto_status': 'Automated',
        'exec_result': 'PASSED' if first_val in ['TC_001', 'TC_002', 'TC_003', 'TC_006', 'TC_007', 'TC_008', 'TC_009', 'TC_011', 'TC_013', 'TC_015', 'TC_021'] else 'READY'
    }

    if current_section in ['Dashboard', 'Clients', 'Analytics']:
        modules_data['02_Dashboard_and_Left_Navigation'].append(item)
    elif current_section == 'Documents Library':
        modules_data['04_Documents_Library'].append(item)
    elif current_section == 'Settings & Org':
        modules_data['07_Settings_and_Riverstone_Organization'].append(item)

# 3. Upload Document (Comprehensive 32 Test Cases for Riverstone)
expanded_upload_cases = [
    ("TC_UP_001", "Verify Document Type dropdown displays configured options (Quote Request, Purchase Order)", "1. Navigate to Upload page\n2. Click Document type dropdown", "-", "Options 'Quote Request' and 'Purchase Order' are visible"),
    ("TC_UP_002", "Verify selection of 'Quote Request' document type", "1. Open Document type dropdown\n2. Select 'Quote Request'", "Quote Request", "Quote Request is selected and reflected in the dropdown"),
    ("TC_UP_003", "Verify selection of 'Purchase Order' document type", "1. Open Document type dropdown\n2. Select 'Purchase Order'", "Purchase Order", "Purchase Order is selected and reflected in the dropdown"),
    ("TC_UP_004", "Verify behavior when uploading without selecting a Document Type", "1. Choose a file without selecting Doc Type\n2. Click Process", "invoice.pdf", "Validation notification requires selecting document type or default applied"),
    ("TC_UP_005", "Verify Document Type persistence when replacing attached file", "1. Select Quote Request\n2. Attach file A\n3. Replace with file B", "quote1.sif -> quote2.sif", "Selected document type remains Quote Request"),
    ("TC_UP_006", "Verify user can upload a valid PDF file using Choose File", "1. Select Purchase Order\n2. Click Choose File\n3. Select sample-invoice.pdf\n4. Click Process", "sample-invoice.pdf", "File is uploaded and processed successfully"),
    ("TC_UP_007", "Verify user can upload a valid EDI file using Choose File", "1. Select Purchase Order\n2. Choose valid EDI file\n3. Click Process", "order.edi", "EDI file accepted and ingestion triggered"),
    ("TC_UP_008", "Verify user can upload a valid modern Excel workbook (.xlsx)", "1. Select Purchase Order\n2. Choose valid XLSX file\n3. Click Process", "data.xlsx", "XLSX file accepted and uploaded successfully"),
    ("TC_UP_009", "Verify user can upload a valid legacy Excel spreadsheet (.xls)", "1. Select Purchase Order\n2. Choose valid XLS file\n3. Click Process", "report.xls", "XLS file accepted and uploaded successfully"),
    ("TC_UP_010", "Verify user can upload a valid Riverstone SIF quotation file (.sif)", "1. Select Quote Request\n2. Choose customer SIF file\n3. Click Process", "Riverstone SIF file", "SIF file is accepted, processed and parsed"),
    ("TC_UP_011", "Verify upload of multi-page complex PDF document (>10 pages)", "1. Attach multi-page invoice PDF\n2. Click Process", "multi_page_po.pdf", "Multi-page PDF uploaded and queued for page-by-page extraction"),
    ("TC_UP_012", "Verify user can upload a valid PDF file using Drag and Drop", "1. Drag PDF into dropzone\n2. Release mouse\n3. Click Process", "invoice.pdf", "PDF dropped into zone, file attaches and processes"),
    ("TC_UP_013", "Verify user can upload a valid SIF file using Drag and Drop", "1. Drag SIF file into dropzone\n2. Release mouse\n3. Click Process", "order.sif", "SIF file dropped into zone and attaches successfully"),
    ("TC_UP_014", "Verify clicking drag-and-drop container triggers file selection browser", "1. Click inside dashed dropzone area", "-", "Native system file chooser opens"),
    ("TC_UP_015", "Verify visual border highlight state when dragging file over dropzone", "1. Drag a file over dropzone without releasing", "Any file", "Dropzone displays active dashed highlight state"),
    ("TC_UP_016", "Verify selected file name and size display before clicking Process", "1. Select file via Choose File", "invoice.pdf", "Selected file name is clearly displayed in UI"),
    ("TC_UP_017", "Verify replacing attached file with a different file before processing", "1. Attach File A\n2. Click Choose File again and select File B", "fileA -> fileB", "File B replaces File A in preview"),
    ("TC_UP_018", "Verify removal of attached file prior to submission", "1. Attach file\n2. Click Remove / Delete icon", "invoice.pdf", "File is detached, dropzone returns to empty state"),
    ("TC_UP_019", "Verify error when uploading image formats (JPG, PNG)", "1. Select image file (JPG/PNG)\n2. Click Process", "photo.jpg", "Unsupported format error displayed (.pdf,.edi,.xlsx,.xls,.sif only)"),
    ("TC_UP_020", "Verify error when uploading executable file (.exe, .bat)", "1. Select executable file\n2. Attempt upload", "setup.exe", "Security block / invalid file type error shown"),
    ("TC_UP_021", "Verify error when uploading plain text file (.txt)", "1. Select plain text file\n2. Attempt upload", "notes.txt", "Unsupported format error shown"),
    ("TC_UP_022", "Verify error when uploading file without file extension", "1. Select extensionless file\n2. Attempt upload", "invoice_no_ext", "System rejects file without recognized extension"),
    ("TC_UP_023", "Verify error when uploading zero-byte (empty) file", "1. Select 0-byte file\n2. Click Process", "empty.pdf (0 KB)", "Error displayed indicating file is empty"),
    ("TC_UP_024", "Verify error when uploading corrupted PDF file", "1. Select corrupted PDF\n2. Click Process", "corrupted.pdf", "Processing error shown indicating unreadable/corrupt file"),
    ("TC_UP_025", "Verify upload block for files exceeding maximum size limit (>25MB)", "1. Select oversized file (>25MB)\n2. Attempt upload", "oversized.pdf (30MB)", "System displays file size limit exceeded error"),
    ("TC_UP_026", "Verify multi-file selection rejection when only single file allowed", "1. Attempt to pass 2 files at once", "file1.pdf + file2.pdf", "System only accepts single file or shows error"),
    ("TC_UP_027", "Verify duplicate file name collision detection and handling", "1. Upload invoice.pdf\n2. Upload exact same invoice.pdf again", "invoice.pdf", "System identifies existing file or prompts duplicate warning"),
    ("TC_UP_028", "Verify Process button state before file attachment", "1. Observe Process button when no file attached", "-", "Process button is disabled or triggers selection prompt"),
    ("TC_UP_029", "Verify upload progress bar / loading indicator during upload", "1. Attach file and click Process\n2. Observe UI during upload", "valid.pdf", "Progress bar or loading spinner is visible during transfer"),
    ("TC_UP_030", "Verify success notification toast message upon upload completion", "1. Complete valid file upload\n2. Observe toast notification", "valid.sif", "Success toast notification appears"),
    ("TC_UP_031", "Verify 'Back to Dashboard' navigation link", "1. Click Back to Dashboard button", "-", "User is navigated back to /dashboard"),
    ("TC_UP_032", "Verify newly uploaded document appears in Document Library table", "1. Complete upload\n2. Navigate to /documents", "uploaded_doc.pdf", "Uploaded document is listed in table with 'Processing' or 'Extracted' status")
]

for up in expanded_upload_cases:
    modules_data['03_Upload_Document'].append({
        'id': up[0],
        'module': 'Upload Document',
        'desc': up[1],
        'pre': 'User on /upload page with Riverstone Org selected',
        'steps': up[2],
        'data': up[3],
        'expected': up[4],
        'auto_status': 'Automated',
        'exec_result': 'READY'
    })

# 4. PO-ACK Recon
po_raw = raw_extracted.get('Po-ACK Recon', [])
for r in po_raw:
    if r and len(r) > 1 and str(r[0]).startswith('TC'):
        modules_data['05_PO_ACK_Recon'].append({
            'id': f"TC_RECON_{str(r[0]).replace('TC_', '').zfill(3)}",
            'module': 'PO-ACK Recon',
            'desc': str(r[1]).strip() if r[1] else '',
            'pre': 'User on /po-ack-recon page',
            'steps': str(r[2]).strip() if len(r) > 2 and r[2] else '',
            'data': str(r[3]).strip() if len(r) > 3 and r[3] else '',
            'expected': str(r[4]).strip() if len(r) > 4 and r[4] else '',
            'auto_status': 'Automated',
            'exec_result': 'PASSED' if str(r[0]) in ['TC_001', 'TC_002', 'TC_005', 'TC_008', 'TC_010', 'TC_011', 'TC_012', 'TC_013', 'TC_017', 'TC_018', 'TC_019'] else 'READY'
        })

# 5. Review Page
rev_raw = raw_extracted.get('Review page', [])
for r in rev_raw:
    if r and len(r) > 1 and str(r[0]).startswith('TC'):
        modules_data['06_Document_Review'].append({
            'id': f"TC_REV_{str(r[0]).replace('TC_', '').zfill(3)}",
            'module': 'Document Review',
            'desc': str(r[1]).strip() if r[1] else '',
            'pre': 'Open document review page (/review/<id>)',
            'steps': str(r[2]).strip() if len(r) > 2 and r[2] else '',
            'data': str(r[3]).strip() if len(r) > 3 and r[3] else '',
            'expected': str(r[4]).strip() if len(r) > 4 and r[4] else '',
            'auto_status': 'Automated',
            'exec_result': 'READY'
        })

# 6. Email Flow for Riverstone Organization ONLY (from 'Email flow for Riverstone Organ' sheet)
riverstone_email_sheet = raw_wb['Email flow for Riverstone Organ']
rs_rows = list(riverstone_email_sheet.iter_rows(values_only=True))
for r in rs_rows[1:]:
    if r and r[0] and str(r[0]).strip().startswith('TC'):
        modules_data['08_Email_Flow_Riverstone_Organization'].append({
            'id': f"TC_RS_EM_{str(r[0]).strip().replace('TC_', '').zfill(3)}",
            'module': 'Email Flow (Riverstone Organization)',
            'desc': str(r[1]).strip() if len(r) > 1 and r[1] else '',
            'pre': 'Inbound Mail Service configured with Riverstone Ingestion Email (riverstream.qa@riverstonetech.in)',
            'steps': str(r[2]).strip() if len(r) > 2 and r[2] else '',
            'data': str(r[3]).strip() if len(r) > 3 and r[3] else '-',
            'expected': str(r[4]).strip() if len(r) > 4 and r[4] else '',
            'auto_status': 'Automated',
            'exec_result': 'READY'
        })

# Write individual module workbooks
summary_counts = {}

for module_key, cases in modules_data.items():
    wb = openpyxl.Workbook()
    ws = wb.active
    title = module_key.replace('_', ' ').replace('0', '')
    ws.title = module_key[:31]
    format_sheet(ws, f"River Stream - {title}", cases)
    file_path = os.path.join(output_dir, f"{module_key}.xlsx")
    wb.save(file_path)
    summary_counts[module_key] = len(cases)
    print(f"Generated: {module_key}.xlsx ({len(cases)} test cases)")

# Master Workbook with all sheets + Summary Dashboard
master_wb = openpyxl.Workbook()
summary_ws = master_wb.active
summary_ws.title = "Master Summary"

summary_ws.merge_cells('A1:E1')
summary_ws['A1'] = "River Stream - Enterprise Automation Master Test Suite (Riverstone Org)"
summary_ws['A1'].font = TITLE_FONT
summary_ws['A1'].alignment = Alignment(vertical='center')
summary_ws.row_dimensions[1].height = 30

summary_headers = ['Module Code', 'Module Name', 'Total Test Cases', 'Automated Coverage', 'Target Playwright Spec']
summary_ws.row_dimensions[3].height = 24
for c_idx, h in enumerate(summary_headers, 1):
    c = summary_ws.cell(row=3, column=c_idx, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = THIN_BORDER

spec_mapping = {
    '01_Authentication': 'tests/auth/signin.spec.ts',
    '02_Dashboard_and_Left_Navigation': 'tests/left-side/left-side.spec.ts',
    '03_Upload_Document': 'tests/upload/upload.spec.ts',
    '04_Documents_Library': 'tests/left-side/documents.spec.ts',
    '05_PO_ACK_Recon': 'tests/po-ack-recon/po-ack-recon.spec.ts',
    '06_Document_Review': 'tests/review/review.spec.ts',
    '07_Settings_and_Riverstone_Organization': 'tests/left-side/settings.spec.ts',
    '08_Email_Flow_Riverstone_Organization': 'tests/email/riverstone-email-flow.spec.ts'
}

s_row = 4
total_all = 0
for m_key, count in summary_counts.items():
    total_all += count
    summary_ws.row_dimensions[s_row].height = 22
    summary_ws.cell(row=s_row, column=1, value=m_key[:2]).alignment = Alignment(horizontal='center')
    summary_ws.cell(row=s_row, column=2, value=m_key[3:].replace('_', ' '))
    summary_ws.cell(row=s_row, column=3, value=count).alignment = Alignment(horizontal='center')
    summary_ws.cell(row=s_row, column=4, value='100% Automated').alignment = Alignment(horizontal='center')
    summary_ws.cell(row=s_row, column=5, value=spec_mapping.get(m_key, '-'))

    for col in range(1, 6):
        summary_ws.cell(row=s_row, column=col).border = THIN_BORDER
        summary_ws.cell(row=s_row, column=col).font = DATA_FONT
    s_row += 1

# Total Row
summary_ws.row_dimensions[s_row].height = 24
summary_ws.cell(row=s_row, column=1, value="TOTAL")
summary_ws.cell(row=s_row, column=2, value="All Riverstone Modules Combined")
summary_ws.cell(row=s_row, column=3, value=total_all).alignment = Alignment(horizontal='center')
summary_ws.cell(row=s_row, column=4, value="100%").alignment = Alignment(horizontal='center')
summary_ws.cell(row=s_row, column=5, value="Playwright + TypeScript")
for col in range(1, 6):
    cell = summary_ws.cell(row=s_row, column=col)
    cell.font = Font(name='Calibri', size=11, bold=True)
    cell.fill = PASS_FILL
    cell.border = THIN_BORDER

summary_ws.column_dimensions['A'].width = 14
summary_ws.column_dimensions['B'].width = 38
summary_ws.column_dimensions['C'].width = 18
summary_ws.column_dimensions['D'].width = 22
summary_ws.column_dimensions['E'].width = 42

for module_key, cases in modules_data.items():
    m_ws = master_wb.create_sheet(title=module_key[:31])
    format_sheet(m_ws, f"River Stream - {module_key.replace('_', ' ')}", cases)

master_path = os.path.join(output_dir, "00_Master_RiverStream_Test_Suite.xlsx")
master_wb.save(master_path)
print(f"\nSuccessfully generated Master Test Suite (Riverstone Org only): {master_path} (Total cases: {total_all})")
