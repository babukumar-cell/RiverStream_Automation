import os
import sys
import zipfile
import json
import xml.etree.ElementTree as ET

def parse_with_openpyxl(file_path):
    import openpyxl
    wb = openpyxl.load_workbook(file_path, data_only=True)
    result = {}
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        result[sheet_name] = rows
    return result

def parse_with_zipfile(file_path):
    ns = {'ns': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(file_path, 'r') as z:
        # Shared strings
        shared_strings = []
        if 'xl/sharedStrings.xml' in z.namelist():
            tree = ET.fromstring(z.read('xl/sharedStrings.xml'))
            for si in tree.findall('ns:si', ns):
                texts = [t.text or '' for t in si.findall('.//ns:t', ns)]
                shared_strings.append(''.join(texts))
        
        # Workbook (sheet names and IDs)
        wb_tree = ET.fromstring(z.read('xl/workbook.xml'))
        sheets_info = []
        for sheet in wb_tree.findall('.//ns:sheet', ns):
            name = sheet.attrib.get('name')
            sheet_id = sheet.attrib.get('sheetId')
            r_id = sheet.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            sheets_info.append({'name': name, 'sheetId': sheet_id, 'rId': r_id})
            
        result = {}
        # Worksheets
        sheet_files = [f for f in z.namelist() if f.startswith('xl/worksheets/sheet') and f.endswith('.xml')]
        sheet_files.sort()
        
        for idx, sfile in enumerate(sheet_files):
            s_name = sheets_info[idx]['name'] if idx < len(sheets_info) else f"Sheet{idx+1}"
            stree = ET.fromstring(z.read(sfile))
            rows_data = []
            for row in stree.findall('.//ns:row', ns):
                row_vals = []
                for c in row.findall('ns:c', ns):
                    val_type = c.attrib.get('t')
                    v_elem = c.find('ns:v', ns)
                    val = v_elem.text if v_elem is not None else ''
                    if val_type == 's' and val.isdigit():
                        idx_s = int(val)
                        val = shared_strings[idx_s] if idx_s < len(shared_strings) else val
                    row_vals.append(val)
                rows_data.append(row_vals)
            result[s_name] = rows_data
        return result

def main():
    file_path = os.path.join(os.path.dirname(__file__), 'River Stream Test cases.xlsx')
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return
        
    print(f"Parsing {file_path}...")
    try:
        data = parse_with_openpyxl(file_path)
        print("Parsed successfully with openpyxl.")
    except Exception as e:
        print(f"openpyxl failed ({e}), falling back to zipfile/xml parser...")
        data = parse_with_zipfile(file_path)
        print("Parsed successfully with zipfile/xml.")

    output_path = os.path.join(os.path.dirname(__file__), 'test_cases_extracted.json')
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        
    print(f"\nSheets found ({len(data)}): {list(data.keys())}")
    for name, rows in data.items():
        print(f"\nSheet '{name}': {len(rows)} rows")
        for r in rows[:5]:
            print("  ", [str(c)[:40] for c in r if str(c).strip()])

if __name__ == '__main__':
    main()
